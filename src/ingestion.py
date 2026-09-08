import io
import json
from pathlib import Path

import pandas as pd
import streamlit as st
from pypdf import PdfReader
from docx import Document

from src.schema import ALIASES
from src.data_utils import best_alias_map, dataframe_to_records_text
from src.analytics import build_analysis
from src.rag_engine import build_faiss_index
from src.gemini_client import is_auth_error
from src.state import invalidate_api_key

SUPPORTED = ["csv", "xlsx", "xls", "json", "txt", "md", "pdf", "docx"]

def _read_file(uploaded):
    ext = Path(uploaded.name).suffix.lower()
    data = uploaded.getvalue()
    tables, texts = {}, []

    if ext == ".csv":
        tables[uploaded.name] = pd.read_csv(io.BytesIO(data))
    elif ext in (".xlsx", ".xls"):
        book = pd.read_excel(io.BytesIO(data), sheet_name=None)
        for sheet, df in book.items():
            tables[f"{uploaded.name}::{sheet}"] = df
    elif ext == ".json":
        obj = json.loads(data.decode("utf-8", errors="ignore"))
        if isinstance(obj, list):
            tables[uploaded.name] = pd.json_normalize(obj)
        elif isinstance(obj, dict):
            # Try tabular JSON first, otherwise keep readable text.
            try:
                tables[uploaded.name] = pd.json_normalize(obj)
            except Exception:
                texts.append((uploaded.name, json.dumps(obj, ensure_ascii=False, indent=2)))
    elif ext in (".txt", ".md"):
        texts.append((uploaded.name, data.decode("utf-8", errors="ignore")))
    elif ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
        texts.append((uploaded.name, text))
    elif ext == ".docx":
        doc = Document(io.BytesIO(data))
        paragraphs = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        texts.append((uploaded.name, paragraphs))
        for t_idx, table in enumerate(doc.tables, 1):
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            if len(rows) >= 2:
                tables[f"{uploaded.name}::table_{t_idx}"] = pd.DataFrame(rows[1:], columns=rows[0])
    else:
        raise ValueError(f"Unsupported file type: {ext or 'unknown'}")
    return tables, texts

def _normalize_tables(raw_tables):
    normalized_tables = {}
    for name, df in raw_tables.items():
        if df is None or df.empty:
            continue
        df = df.copy()
        df.columns = [str(c).strip() for c in df.columns]
        mapping = best_alias_map(df.columns, ALIASES)
        df = df.rename(columns=mapping)
        normalized_tables[name] = df
    return normalized_tables

def _quality_report(tables):
    columns = set()
    for df in tables.values():
        columns.update(df.columns)

    checks = {
        "Quarter target": any(c in columns for c in ("quarter_target", "quarter")),
        "Production line": "line_id" in columns,
        "Line supervisor": "supervisor" in columns,
        "Weekly target": "target_units" in columns,
        "Actual production": "actual_units" in columns,
        "Machine ID": "machine_id" in columns,
        "Machine issue": "issue" in columns,
        "Downtime": ("downtime_minutes" in columns) or ("downtime_hours" in columns),
        "Defects / rejects": "defect_units" in columns,
    }
    found = [k for k, v in checks.items() if v]
    missing = [k for k, v in checks.items() if not v]
    return {"checks": checks, "found": found, "missing": missing, "score": round(100 * len(found) / len(checks))}

def _make_chunks(tables, texts):
    chunks = []
    metadata = []
    for name, df in tables.items():
        for text in dataframe_to_records_text(name, df):
            chunks.append(text)
            metadata.append({"source": name, "type": "table_row"})
    for name, text in texts:
        clean = " ".join(text.split())
        size = 1100
        overlap = 150
        start = 0
        while start < len(clean):
            chunk = clean[start:start+size]
            if chunk.strip():
                chunks.append(f"Source document {name}. {chunk}")
                metadata.append({"source": name, "type": "document_text"})
            if start + size >= len(clean):
                break
            start += size - overlap
    return chunks, metadata

def render_upload_and_process():
    st.markdown("## Upload factory data")
    st.caption("Supported: CSV, Excel, JSON, TXT, Markdown, PDF and DOCX. Multiple files can be combined in one session.")

    uploads = st.file_uploader(
        "Drop your production, target, line, supervisor and machine-log files here",
        type=SUPPORTED,
        accept_multiple_files=True,
    )

    if uploads:
        st.write(f"**{len(uploads)} file(s) selected**")
        for f in uploads:
            st.caption(f"• {f.name} — {f.size / 1024:.1f} KB")

        if st.button("Analyze & build knowledge index", type="primary", use_container_width=True):
            try:
                raw_tables, raw_texts, file_summary = {}, [], []
                with st.status("Processing factory data…", expanded=True) as status:
                    for up in uploads:
                        st.write(f"Reading **{up.name}**")
                        tables, texts = _read_file(up)
                        raw_tables.update(tables)
                        raw_texts.extend(texts)
                        file_summary.append({"name": up.name, "tables": len(tables), "text_docs": len(texts)})

                    st.write("Normalizing fields")
                    normalized = _normalize_tables(raw_tables)

                    st.write("Checking data sufficiency")
                    quality = _quality_report(normalized)

                    st.write("Calculating production and machine analytics")
                    analysis = build_analysis(normalized)

                    st.write("Creating chunks and Gemini embeddings")
                    chunks, metadata = _make_chunks(normalized, raw_texts)
                    if not chunks:
                        raise ValueError("No readable records were found in the uploaded files.")

                    index = build_faiss_index(st.session_state.client, chunks)

                    st.session_state.raw_tables = raw_tables
                    st.session_state.raw_texts = raw_texts
                    st.session_state.normalized = normalized
                    st.session_state.quality_report = quality
                    st.session_state.analysis = analysis
                    st.session_state.chunks = chunks
                    st.session_state.chunk_metadata = metadata
                    st.session_state.faiss_index = index
                    st.session_state.file_summary = file_summary
                    st.session_state.data_ready = True
                    status.update(label="Analysis ready", state="complete", expanded=False)

                st.success("Data analyzed and FAISS knowledge index built.")
                st.balloons()

            except Exception as exc:
                if is_auth_error(exc):
                    invalidate_api_key(f"Gemini authentication failed: {exc}")
                    st.rerun()
                st.error(f"Could not process the data: {exc}")

    if st.session_state.data_ready:
        st.divider()
        q = st.session_state.quality_report
        c1, c2 = st.columns([1, 1])
        with c1:
            st.metric("Data readiness", f"{q.get('score', 0)}%")
        with c2:
            st.metric("Searchable chunks", f"{len(st.session_state.chunks):,}")
        with st.expander("Detected and missing information", expanded=True):
            for label, ok in q.get("checks", {}).items():
                st.write(("✅" if ok else "⚠️") + " " + label)
